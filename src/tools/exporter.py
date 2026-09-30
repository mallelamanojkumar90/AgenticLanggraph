import io
import re
import html
import logging
from pathlib import Path
from typing import Optional, Union, List

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

import docx
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

def markdown_to_pdf_bytes(markdown_text: str) -> bytes:
    """Converts GitHub-flavored Markdown text into styled PDF bytes using ReportLab."""
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        alignment=0, # Left-aligned
        spaceAfter=14
    )
    
    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#334155'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=6
    )
    
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=4
    )
    
    story = []
    
    for raw_line in markdown_text.split('\n'):
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 4))
            continue
            
        if line.startswith('---'):
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#CBD5E1'), spaceAfter=8))
            continue
            
        if line.startswith('# '):
            story.append(Paragraph(html.escape(line[2:].strip()), title_style))
        elif line.startswith('## '):
            story.append(Paragraph(html.escape(line[3:].strip()), h1_style))
        elif line.startswith('### '):
            story.append(Paragraph(html.escape(line[4:].strip()), h2_style))
        elif line.startswith(('- ', '* ')):
            # Handle bullet item with inline bold
            cleaned_bullet = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', html.escape(line[2:].strip()))
            story.append(Paragraph(f"&bull;&nbsp; {cleaned_bullet}", bullet_style))
        elif re.match(r'^\d+\.\s', line):
            # Numbered list
            cleaned_num = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', html.escape(line.strip()))
            story.append(Paragraph(cleaned_num, bullet_style))
        else:
            cleaned_para = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', html.escape(line))
            # Format markdown links like [Title](url) to clickable blue links
            cleaned_para = re.sub(r'\[([^\]]+)\]\((https?://[^\)]+)\)', r'<a href="\2" color="#0284C7"><u>\1</u></a>', cleaned_para)
            story.append(Paragraph(cleaned_para, body_style))
            
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    doc.build(story)
    return buf.getvalue()

def markdown_to_docx_bytes(markdown_text: str) -> bytes:
    """Converts Markdown text into a styled Microsoft Word .docx document bytes."""
    doc = docx.Document()
    
    for raw_line in markdown_text.split('\n'):
        line = raw_line.strip()
        if not line:
            continue
            
        if line.startswith('---'):
            continue
            
        if line.startswith('# '):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.style.font.name = 'Calibri'
            h.style.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith('## '):
            h = doc.add_heading(line[3:].strip(), level=2)
            h.style.font.name = 'Calibri'
            h.style.font.color.rgb = RGBColor(30, 41, 59)
        elif line.startswith('### '):
            h = doc.add_heading(line[4:].strip(), level=3)
            h.style.font.name = 'Calibri'
            h.style.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith(('- ', '* ')):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:].strip()
            _add_markdown_runs(p, text)
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s*', '', line).strip()
            _add_markdown_runs(p, text)
        else:
            p = doc.add_paragraph()
            _add_markdown_runs(p, line)
            
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def _add_markdown_runs(paragraph, text: str):
    """Splits markdown inline bold (**text**) and plain text into paragraph runs."""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

def generate_mermaid_diagram(topic: str, sections_data: List[dict]) -> str:
    """Generates a Mermaid graph diagram representing the research architecture and subtopics."""
    clean_topic = re.sub(r'["\n]', '', topic)[:35]
    lines = [
        "```mermaid",
        "flowchart TD",
        f'    Root["{clean_topic}"]'
    ]
    for idx, s in enumerate(sections_data, 1):
        title = re.sub(r'["\n]', '', s.get("section_title", f"Section {idx}"))[:30]
        node_id = f"Sec{idx}"
        num_sources = len(s.get("sources", []))
        lines.append(f'    Root --> {node_id}["{title}<br/>({num_sources} sources)"]')
    lines.append("```")
    return "\n".join(lines)
